"""Exportación del informe a Excel (openpyxl)."""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from core import TIPOS_DOC_EXCLUIDOS, motivo_exclusion, Base, CuentaRet, ResTarifa, calc_autorretencion, totales_retenciones, GRUPOS_RETENCION

NUM = "#,##0.00"
TIPO = {"practicada": "Retención practicada", "autorretencion": "Autorretención contabilizada", "pago": "Pago / cierre"}


def _f(x):
    return float(x) if isinstance(x, Decimal) else x


def _sheet(wb, title, rows, widths, num_cols=()):
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append([_f(x) for x in r])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in num_cols:
        for cell in ws[col][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = NUM
    return ws


def _hoja_ica(wb, cuentas: list[CuentaRet], invertir: bool) -> None:
    """Hoja 'ICA 2368 por C. costos': subtotales por centro de costos (resumen) + detalle de movimientos."""
    sgn = -1 if invertir else 1
    SIN = "(Sin centro de costos)"
    centros: dict[str, dict] = {}
    for c in cuentas:
        for m in c.movs:
            cc = (m.get("centrocosto") or "").strip() or SIN
            d = centros.setdefault(cc, {"cuentas": {}, "movs": []})
            a = d["cuentas"].setdefault(c.codigo, dict(nombre=c.nombre, base=Decimal(0), causada=Decimal(0), dev=Decimal(0)))
            if m["valor"] < 0:
                a["causada"] += -m["valor"]
                a["base"] += m.get("baseimpuesto", Decimal(0))
            elif m["valor"] > 0:
                a["dev"] += m["valor"]
            d["movs"].append((c, m))
    orden = sorted(centros, key=lambda k: (k == SIN, k))
    ws = wb.create_sheet("ICA 2368 por C. costos")
    bold = Font(bold=True)
    ws.append(["RETENCIÓN DE ICA (CUENTAS 2368) — SUBTOTAL POR CENTRO DE COSTOS"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Devoluciones / reversiones en negativo; el neto = causada + devoluciones."])
    ws.append([])
    ws.append(["Centro de costos", "Cuenta", "Nombre cuenta", "Base (movimientos de causación)", "Retención causada",
               "Devoluciones / reversiones (−)", "Neto"])
    for c in ws[ws.max_row]:
        c.font = bold
    tot = [Decimal(0)] * 4
    for cc in orden:
        sub = [Decimal(0)] * 4
        for cod, a in sorted(centros[cc]["cuentas"].items()):
            fila = [a["base"], a["causada"], -a["dev"], a["causada"] - a["dev"]]
            ws.append([cc, cod, a["nombre"], *[float(v) for v in fila]])
            sub = [s + v for s, v in zip(sub, fila)]
        ws.append([f"Subtotal {cc}", "", "", *[float(v) for v in sub]])
        for c in ws[ws.max_row]:
            c.font = bold
        tot = [s + v for s, v in zip(tot, sub)]
    ws.append(["TOTAL RETENCIÓN DE ICA (2368)", "", "", *[float(v) for v in tot]])
    for c in ws[ws.max_row]:
        c.font = bold
    resumen_fin = ws.max_row
    ws.append([])
    ws.append(["DETALLE DE MOVIMIENTOS POR CENTRO DE COSTOS"]); ws[ws.max_row][0].font = bold
    ws.append(["Centro de costos", "Cuenta", "Nombre cuenta", "Fecha", "Doc.", "Tercero", "Identificación", "Descripción",
               "Base impuesto", "% ", "Valor"])
    for c in ws[ws.max_row]:
        c.font = bold
    for cc in orden:
        s = Decimal(0)
        for c, m in sorted(centros[cc]["movs"], key=lambda t: (t[1]["fecha"], t[0].codigo)):
            v = m["valor"] * sgn
            s += v
            ws.append([cc, c.codigo, c.nombre, m["fecha"], f"{m['tipodocumento']} {m['nrodocumento']}".strip(), m["tercero"],
                       m["identificacion"], m["descripcion"], _f(m.get("baseimpuesto")), _f(m.get("porcentajeimpuesto")), float(v)])
        ws.append([f"Subtotal {cc}", "", "", "", "", "", "", "", "", "", float(s)])
        for c in ws[ws.max_row]:
            c.font = bold
    for i, w in enumerate([34, 14, 38, 24, 22, 28, 18, 50, 18, 8, 20], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = NUM


def exportar(empresa: str, desde: str, hasta: str, base: Base, tr: ResTarifa, checks: list[dict],
             retenciones: list[CuentaRet] | None = None, excluidos: list[dict] | None = None,
             nits: list[str] | None = None) -> bytes:
    ret = retenciones or []
    exc = excluidos or []
    ok = tr.status == "ok"
    pct = tr.tarifa if ok else None
    auto = calc_autorretencion(base.base, pct) if ok else None
    tot = totales_retenciones(ret)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumen")
    filas = [
        ["LIQUIDACIÓN PRELIMINAR DE AUTORRETENCIÓN — sujeta a revisión"], [],
        ["Empresa", empresa], ["Período de liquidación", f"{desde} a {hasta}"],
        ["Tarifa de autorretención (Maestro de empresas) %", _f(pct) if ok else "No configurada"],
        ["Signo de presentación", "Invertido (ingresos positivos)" if base.invertir_signo else "Tal cual el auxiliar"], [],
        ["Base total de autorretención", _f(base.base)],
        ["Porcentaje aplicado (%)", _f(pct) if ok else None],
        ["Autorretención calculada", _f(auto) if ok else "No calculada"], [],
        ["RETENCIONES PRACTICADAS EN EL PERÍODO (cuentas 2365 y 2368)"],
        ["2365 Retención en la fuente — causada (créditos)", _f(tot["2365"]["creditos"])],
        ["2365 Retención en la fuente — devoluciones / reversiones (−)", -_f(tot["2365"]["debitos"])],
        ["2368 Retención de ICA — causada (créditos)", _f(tot["2368"]["creditos"])],
        ["2368 Retención de ICA — devoluciones / reversiones (−)", -_f(tot["2368"]["debitos"])],
        ["Total retenciones practicadas (neto de devoluciones)", _f(tot["total"]["creditos"] - tot["total"]["debitos"])],
        ["NIT excluidos del cálculo", ", ".join(nits or []) or "Ninguno"],
        ["Movimientos excluidos (NIT y tipo de documento " + ", ".join(TIPOS_DOC_EXCLUIDOS) + ")", len(exc)], [],
        ["La parametrización de la tarifa no sustituye la validación de las reglas tributarias aplicables a la empresa y al período."],
    ]
    for f in filas:
        ws.append(f)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 58, 46
    for r in (8, 10, 13, 14, 15, 16, 17):
        if isinstance(ws[f"B{r}"].value, float):
            ws[f"B{r}"].number_format = NUM
    for r in (8, 10, 12, 17):
        ws[f"A{r}"].font = Font(bold=True)

    det = [["Código", "Cuenta", "Movimientos", "Suma valor (auxiliar)", "Valor neto", "Autorretención"]]
    det += [[c.codigo, c.nombre, c.n, c.raw, c.valor, calc_autorretencion(c.valor, pct) if ok else None] for c in base.cuentas]
    det.append(["TOTAL", "Base total", base.n_movs, base.base_raw, base.base, auto])
    w = _sheet(wb, "Detalle por cuenta", det, [14, 40, 12, 22, 20, 18], "DEF")
    for c in w[w.max_row]:
        c.font = Font(bold=True)

    # ---- retenciones 2365 / 2368
    rt = [["Grupo", "Código", "Cuenta", "Tipo de cuenta", "Tarifa(s) %", "Base (movimientos de causación)",
           "Retención causada (créditos)", "Devoluciones / reversiones / débitos (−)", "Neto", "Movimientos"]]
    for g, nombre in GRUPOS_RETENCION.items():
        cs = [c for c in ret if c.grupo == g]
        for c in cs:
            rt.append([f"{g} {nombre}", c.codigo, c.nombre, TIPO[c.tipo], ", ".join(f"{t:f}" for t in sorted(c.tarifas)),
                       c.base, c.creditos, -c.debitos, c.creditos - c.debitos, c.n])
        t = tot[g]
        rt.append([f"Subtotal {g} (solo retenciones practicadas)", "", "", "", "", t["base"], t["creditos"], -t["debitos"], t["creditos"] - t["debitos"], ""])
    t = tot["total"]
    rt.append(["TOTAL RETENCIONES PRACTICADAS (neto)", "", "", "", "", t["base"], t["creditos"], -t["debitos"], t["creditos"] - t["debitos"], ""])
    w = _sheet(wb, "Retenciones 2365-2368", rt, [44, 14, 40, 26, 12, 24, 24, 26, 20, 12], "FGHI")
    for row in w.iter_rows(min_row=2):
        if str(row[0].value).startswith(("Subtotal", "TOTAL")):
            for c in row:
                c.font = Font(bold=True)

    pp = [["Grupo", "Código", "Cuenta", "Tipo de persona", "Base (causación)", "Retención causada",
           "Reversiones / devoluciones (−)", "Neto"]]
    por_tipo: dict[str, list] = {}
    for c in ret:
        if c.tipo == "practicada":
            for per, (b, r, d) in sorted(c.por_persona.items()):
                pp.append([c.grupo, c.codigo, c.nombre, per, b, r, -d, r - d])
                t = por_tipo.setdefault(per, [Decimal(0)] * 4)
                por_tipo[per] = [t[0] + b, t[1] + r, t[2] - d, t[3] + r - d]
    pp.append([])
    pp.append(["RESUMEN POR TIPO DE PERSONA"])
    ini_resumen = len(pp)
    for per, v in sorted(por_tipo.items()):
        pp.append(["Total", "", "", per, *v])
    if por_tipo:
        pp.append(["TOTAL GENERAL", "", "", "", *[sum(v[i] for v in por_tipo.values()) for i in range(4)]])
    w = _sheet(wb, "Retenciones por tipo persona", pp, [10, 14, 44, 18, 22, 22, 28, 20], "EFGH")
    w[ini_resumen][0].font = Font(bold=True)
    for row in w.iter_rows(min_row=ini_resumen + 1):
        if row[0].value in ("Total", "TOTAL GENERAL"):
            for cell in row:
                cell.font = Font(bold=True)

    ica = [c for c in ret if c.grupo == "2368"]
    if ica:
        _hoja_ica(wb, ica, base.invertir_signo)

    mr = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación", "Tercero",
           "Tipo persona", "Descripción", "Base impuesto", "% impuesto", "valor"]]
    for c in ret:
        mr += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
                m["identificacion"], m["tercero"], m.get("tipo_persona", ""), m["descripcion"], m.get("baseimpuesto"),
                m.get("porcentajeimpuesto"), m["valor"]] for m in c.movs]
    _sheet(wb, "Mov. retenciones", mr, [12, 30, 10, 10, 12, 8, 10, 14, 30, 12, 40, 18, 10, 18], "LN")

    mv = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación",
           "Tercero", "Descripción", "Centro costo", "valor"]]
    for c in base.cuentas:
        mv += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
                m["identificacion"], m["tercero"], m["descripcion"], m["centrocosto"], m["valor"]] for m in c.movs]
    _sheet(wb, "Mov. ingresos", mv, [12, 30, 10, 10, 12, 8, 10, 14, 30, 40, 20, 18], "L")

    ex = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación", "Tercero",
           "Descripción", "valor", "Motivo de exclusión"]]
    ex += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
            m["identificacion"], m["tercero"], m["descripcion"], m["valor"], motivo_exclusion(m, nits or [])] for m in exc]
    _sheet(wb, "Excluidos (NIT y CDI)", ex, [12, 40, 10, 10, 12, 8, 10, 14, 24, 60, 18, 28], "K")

    ng = [["Código", "Cuenta", "Valor neto"]] + [[c.codigo, c.nombre, c.valor] for c in base.negativas]
    ng += [[], ["Efecto de las negativas en la base", "", base.efecto_negativas], ["Base sin cuentas negativas", "", base.base_sin_negativas]]
    _sheet(wb, "Cuentas negativas", ng, [34, 40, 20], "C")

    vl = [["Control", "Resultado", "Detalle"]] + [[c["titulo"], "OK" if c["ok"] else ("Revisar" if c["nivel"] == "warn" else "Falla"), c["detalle"]] for c in checks]
    _sheet(wb, "Validaciones", vl, [48, 10, 100])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
