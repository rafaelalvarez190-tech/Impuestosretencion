"""Núcleo de cálculo de la autorretención (sin dependencias de Streamlit).

Todo el dinero se maneja con Decimal. La base sale EXCLUSIVAMENTE de la columna `valor`.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Iterable

PREFIJO_INGRESOS = "4"
CENT = Decimal("0.01")
REQUIRED = ["empresa", "fecha", "codigocuenta", "nombrecuenta", "valor"]


# ---------------------------------------------------------------- números y fechas
def parse_number(v: Any) -> Decimal:
    """Acepta Decimal/int/float o texto es-CO ('1.234,56', '182024,87', '(100,00)')."""
    if v is None or v == "":
        return Decimal(0)
    if isinstance(v, Decimal):
        return v
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = re.sub(r"[\s$]", "", str(v))
    if not s:
        return Decimal(0)
    neg = s.startswith("(") and s.endswith(")")
    if neg:
        s = s[1:-1]
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    d = Decimal(s)  # lanza InvalidOperation si no es número
    return -d if neg else d


def q2(d: Decimal) -> Decimal:
    return d.quantize(CENT, rounding=ROUND_HALF_UP)


def normalize_date(v: Any) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):  # serial de Excel
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m[1]}-{m[2]}-{m[3]}"
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        return f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}"
    return ""


# ---------------------------------------------------------------- lectura de archivos
def read_table(data: bytes, filename: str) -> list[list[Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)
                if any(c not in (None, "") for c in r)]
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")
    first = text.split("\n", 1)[0]
    delim = ";" if first.count(";") > first.count(",") else ","
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]


def _norm_header(h: Any) -> str:  # "tipo_persona" -> "tipopersona"
    s = unicodedata.normalize("NFD", str(h).strip().lower())
    return re.sub(r"[\s_]+", "", "".join(c for c in s if unicodedata.category(c) != "Mn"))


def normalize_rows(table: list[list[Any]]) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    if not table:
        return [], ["El archivo está vacío."]
    headers = [_norm_header(h) for h in table[0]]
    missing = [h for h in REQUIRED if h not in headers]
    if missing:
        return [], [f"Faltan columnas obligatorias: {', '.join(missing)}."]
    ix = {h: i for i, h in enumerate(headers)}

    def get(r, h):
        i = ix.get(h)
        return "" if i is None or i >= len(r) or r[i] is None else r[i]

    rows = []
    for n, r in enumerate(table[1:], start=2):
        codigo, empresa = str(get(r, "codigocuenta")).strip(), str(get(r, "empresa")).strip()
        fecha = normalize_date(get(r, "fecha"))
        if not codigo or not empresa:
            errors.append(f"Fila {n}: sin empresa o código de cuenta.")
            continue
        if not fecha:
            errors.append(f"Fila {n}: fecha inválida ({get(r, 'fecha')}).")
            continue
        try:
            valor = parse_number(get(r, "valor"))
            debito, credito = parse_number(get(r, "debito")), parse_number(get(r, "credito"))
            base_imp = parse_number(get(r, "baseimpuesto"))
            pct_imp = parse_number(get(r, "porcentajeimpuesto"))
        except Exception:
            errors.append(f"Fila {n}: valor no numérico ({get(r, 'valor')}).")
            continue
        s = lambda h: str(get(r, h)).strip()  # noqa: E731
        m = dict(
            id=s("id"), empresa=empresa, centrocosto=s("centrocosto"), fecha=fecha,
            tipodocumento=s("tipodocumento"), nrodocumento=s("nrodocumento"),
            codigocuenta=codigo, nombrecuenta=s("nombrecuenta"), descripcion=s("descripcion"),
            factura=s("factura"), identificacion=s("identificacion"), tercero=s("tercero"),
            valor=valor, debito=debito, credito=credito, fila=n,  # débito/crédito solo informativos
            baseimpuesto=base_imp, porcentajeimpuesto=pct_imp, tipo_persona=s("tipopersona"),
        )
        m["key"] = f"{empresa}|{m['id']}" if m["id"] else f"{empresa}|h:{fingerprint(m)}"
        rows.append(m)
    return rows, errors


def fingerprint(m: dict) -> str:
    raw = "¦".join(str(m[k]) for k in ("empresa", "fecha", "tipodocumento", "nrodocumento",
                                      "codigocuenta", "identificacion", "descripcion", "factura", "valor"))
    return hashlib.sha1(raw.encode()).hexdigest()[:16]


# ---------------------------------------------------------------- NIT excluidos
NITS_EXCLUIDOS_DEFECTO = "800197268"


def solo_digitos(x: Any) -> str:
    return re.sub(r"\D", "", str(x or ""))


def parse_nits(texto: str) -> list[str]:
    """'800197268, 900.123.456-1' -> ['800197268', '9001234561']"""
    return [d for d in (solo_digitos(p) for p in re.split(r"[,;\n]", texto or "")) if d]


def nit_excluido(identificacion: Any, nits: list[str]) -> bool:
    """Coincide con el NIT sin DV (800197268) o con DV pegado (8001972684)."""
    d = solo_digitos(identificacion)
    return bool(d) and any(d == n or (len(d) == len(n) + 1 and d.startswith(n)) for n in nits)


TIPOS_DOC_EXCLUIDOS = ("CDI",)


def motivo_exclusion(m: dict, nits: list[str]) -> str:
    """'' si el movimiento entra al cálculo; si no, el motivo (NIT excluido o tipo de documento excluido)."""
    if str(m.get("tipodocumento") or "").strip().upper() in TIPOS_DOC_EXCLUIDOS:
        return "Tipo de documento " + str(m["tipodocumento"]).strip().upper()
    if nit_excluido(m.get("identificacion"), nits):
        return "NIT excluido"
    return ""


def separar_excluidos(movs: Iterable[dict], nits: list[str]) -> tuple[list[dict], list[dict]]:
    """(incluidos, excluidos): no entran a ningún cálculo los movimientos con NIT excluido ni los de tipo de documento CDI."""
    inc, exc = [], []
    for m in movs:
        (exc if motivo_exclusion(m, nits) else inc).append(m)
    return inc, exc


def es_cuenta_ingreso(codigo: str) -> bool:
    return str(codigo).strip().startswith(PREFIJO_INGRESOS)


# ---------------------------------------------------------------- base
@dataclass
class Cuenta:
    codigo: str
    nombre: str
    raw: Decimal = Decimal(0)       # suma de `valor` tal cual el auxiliar
    valor: Decimal = Decimal(0)     # valor presentado (con o sin inversión de signo)
    movs: list[dict] = field(default_factory=list)

    @property
    def n(self) -> int:
        return len(self.movs)


@dataclass
class Base:
    cuentas: list[Cuenta]
    invertir_signo: bool
    excluidos_no4: int

    @property
    def base(self) -> Decimal:
        return sum((c.valor for c in self.cuentas), Decimal(0))

    @property
    def base_raw(self) -> Decimal:
        return sum((c.raw for c in self.cuentas), Decimal(0))

    @property
    def n_movs(self) -> int:
        return sum(c.n for c in self.cuentas)

    @property
    def negativas(self) -> list[Cuenta]:
        return [c for c in self.cuentas if c.valor < 0]

    @property
    def efecto_negativas(self) -> Decimal:
        return sum((c.valor for c in self.negativas), Decimal(0))

    @property
    def base_sin_negativas(self) -> Decimal:
        return self.base - self.efecto_negativas


def build_base(movs: Iterable[dict], invertir_signo: bool = True) -> Base:
    """Base = suma neta (con signo) de `valor` de las cuentas que empiezan por 4."""
    acc: dict[str, Cuenta] = {}
    excl = 0
    for m in movs:
        if not es_cuenta_ingreso(m["codigocuenta"]):
            excl += 1
            continue
        c = acc.setdefault(m["codigocuenta"], Cuenta(m["codigocuenta"], m["nombrecuenta"]))
        c.raw += m["valor"]
        c.movs.append(m)
    sgn = -1 if invertir_signo else 1
    for c in acc.values():
        c.valor = c.raw * sgn
    return Base(sorted(acc.values(), key=lambda c: c.codigo), invertir_signo, excl)


def filtrar(movs: Iterable[dict], empresa: str, desde: str, hasta: str) -> list[dict]:
    return [m for m in movs if m["empresa"] == empresa and desde <= m["fecha"] <= hasta]


def calc_autorretencion(base: Decimal, porcentaje: Decimal | float | str) -> Decimal:
    return q2(base * Decimal(str(porcentaje)) / Decimal(100))


# ---------------------------------------------------------------- maestro de empresas
def validar_empresa(e: dict, existentes: list[str] | None = None) -> list[str]:
    errs = []
    if not str(e.get("nombre", "")).strip():
        errs.append("El nombre de la empresa es obligatorio.")
    if existentes and str(e.get("nombre", "")).strip() in existentes:
        errs.append(f"La empresa {e['nombre']} está repetida.")
    t = e.get("tarifa")
    if t in (None, "") or (isinstance(t, float) and t != t):
        errs.append(f"{e.get('nombre') or 'Empresa'}: falta la tarifa de autorretención.")
    else:
        try:
            if not (0 <= Decimal(str(t)) <= 100):
                errs.append(f"{e.get('nombre')}: la tarifa debe estar entre 0 y 100.")
        except Exception:
            errs.append(f"{e.get('nombre')}: tarifa no numérica.")
    return errs


@dataclass
class ResTarifa:
    status: str  # ok | none | inactive
    tarifa: Decimal | None = None
    msg: str = ""


def resolver_tarifa(maestro: list[dict], empresa: str) -> ResTarifa:
    """La tarifa sale del maestro de empresas. Sin tarifa (o empresa inactiva) no se calcula."""
    e = next((x for x in maestro if x["nombre"] == empresa), None)
    if e is None or e.get("tarifa") in (None, ""):
        return ResTarifa("none", None, f"La empresa {empresa} no tiene tarifa de autorretención en el Maestro de empresas. "
                                       "Configúrela para calcular.")
    if not e.get("activa", True):
        return ResTarifa("inactive", None, f"La empresa {empresa} está inactiva en el Maestro de empresas.")
    return ResTarifa("ok", Decimal(str(e["tarifa"])))


# ---------------------------------------------------------------- retenciones 2365 / 2368
GRUPOS_RETENCION = {"2365": "Retención en la fuente", "2368": "Retención de ICA (industria y comercio)"}


def grupo_retencion(codigo: str) -> str | None:
    c = str(codigo).strip()
    return next((g for g in GRUPOS_RETENCION if c.startswith(g)), None)


@dataclass
class CuentaRet:
    codigo: str
    nombre: str
    grupo: str
    tipo: str                                   # practicada | autorretencion | pago
    creditos: Decimal = Decimal(0)              # retención causada (movimientos con valor < 0), en magnitud
    debitos: Decimal = Decimal(0)               # cierres, pagos, devoluciones (valor > 0), en magnitud
    base: Decimal = Decimal(0)                  # suma de baseimpuesto de los movimientos de causación
    tarifas: set = field(default_factory=set)
    por_persona: dict = field(default_factory=dict)   # tipo_persona -> [base causación, retención causada, devoluciones]
    movs: list[dict] = field(default_factory=list)
    invertir: bool = True

    @property
    def neto(self) -> Decimal:
        n = self.creditos - self.debitos
        return n if self.invertir else -n

    @property
    def n(self) -> int:
        return len(self.movs)


def _tipo_cuenta(nombre: str) -> str:
    u = nombre.upper()
    if "AUTORRETEN" in u:
        return "autorretencion"
    if "PAGO" in u:
        return "pago"
    return "practicada"


def build_retenciones(movs: Iterable[dict], invertir_signo: bool = True) -> list[CuentaRet]:
    """Cuentas 2365 y 2368 del período. Causación = movimientos con valor < 0 (crédito); el resto son
    cierres/pagos/devoluciones. La base y el % salen de baseimpuesto / porcentajeimpuesto del auxiliar."""
    acc: dict[str, CuentaRet] = {}
    for m in movs:
        g = grupo_retencion(m["codigocuenta"])
        if not g:
            continue
        c = acc.setdefault(m["codigocuenta"], CuentaRet(m["codigocuenta"], m["nombrecuenta"], g,
                                                         _tipo_cuenta(m["nombrecuenta"]), invertir=invertir_signo))
        c.movs.append(m)
        v = m["valor"]
        if v < 0:
            c.creditos += -v
            c.base += m.get("baseimpuesto", Decimal(0))
            p = c.por_persona.setdefault(m.get("tipo_persona") or "Sin clasificar", [Decimal(0), Decimal(0), Decimal(0)])
            p[0] += m.get("baseimpuesto", Decimal(0))
            p[1] += -v
        elif v > 0:
            c.debitos += v
            p = c.por_persona.setdefault(m.get("tipo_persona") or "Sin clasificar", [Decimal(0), Decimal(0), Decimal(0)])
            p[2] += v
        if m.get("porcentajeimpuesto"):
            c.tarifas.add(m["porcentajeimpuesto"].normalize())
    return sorted(acc.values(), key=lambda c: c.codigo)


def totales_retenciones(cuentas: list[CuentaRet]) -> dict:
    """Totales de retenciones PRACTICADAS (excluye cuentas de autorretención y de pago) por grupo."""
    out = {}
    for g, nombre in GRUPOS_RETENCION.items():
        cs = [c for c in cuentas if c.grupo == g and c.tipo == "practicada"]
        out[g] = dict(nombre=nombre, base=sum((c.base for c in cs), Decimal(0)),
                      creditos=sum((c.creditos for c in cs), Decimal(0)),
                      debitos=sum((c.debitos for c in cs), Decimal(0)),
                      neto=sum((c.neto for c in cs), Decimal(0)))
    out["total"] = {k: sum((out[g][k] for g in GRUPOS_RETENCION), Decimal(0)) for k in ("base", "creditos", "debitos", "neto")}
    return out


# ---------------------------------------------------------------- controles
def fmt(n: Decimal | float, dec: int = 2) -> str:
    s = f"{Decimal(n):,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def ejecutar_controles(movs_periodo: list[dict], base: Base, tr: ResTarifa, total_movs: int, claves_unicas: int,
                       retenciones: list[CuentaRet] | None = None, auto: Decimal | None = None,
                       excluidos: list[dict] | None = None, nits: list[str] | None = None) -> list[dict]:
    out = []

    def add(titulo, ok, detalle, nivel="error"):
        out.append(dict(titulo=titulo, ok=ok, detalle=detalle, nivel="ok" if ok else nivel))

    no4 = [c.codigo for c in base.cuentas if not es_cuenta_ingreso(c.codigo)]
    add("Todas las cuentas incluidas comienzan por 4", not no4,
        f"Fuera de regla: {', '.join(no4)}" if no4 else
        f"{len(base.cuentas)} cuenta(s) de ingresos; {base.excluidos_no4} movimiento(s) de otras clases excluidos.")
    add("Base = suma neta de la columna `valor`", sum((m['valor'] for c in base.cuentas for m in c.movs), Decimal(0)) == base.base_raw,
        f"Suma de `valor` (signo del auxiliar): {fmt(base.base_raw)}. No se usan débito, crédito ni saldos.")
    neg = base.negativas
    add("Cuentas con saldo negativo revisadas", not neg,
        f"{len(neg)} cuenta(s) reducen la base en {fmt(abs(base.efecto_negativas))}. Base sin ellas: {fmt(base.base_sin_negativas)}."
        if neg else "Ninguna cuenta de ingresos con saldo negativo.", "warn")
    add("Tarifa tomada del Maestro de empresas", tr.status == "ok",
        f"Tarifa de autorretención: {tr.tarifa}%." if tr.status == "ok" else tr.msg)
    add("Hay movimientos de la empresa en el período", bool(movs_periodo),
        f"{len(movs_periodo)} movimiento(s) en el rango." if movs_periodo else "No hay movimientos en el rango elegido.")
    add("Sin movimientos duplicados", total_movs == claves_unicas,
        f"{total_movs} movimiento(s) únicos almacenados (clave: empresa + id del auxiliar).")
    add("Base rastreable hasta los movimientos originales", True,
        "Cada cuenta enlaza a sus movimientos (fila del archivo e id del auxiliar).")
    exc = excluidos or []
    add("Movimientos excluidos (NIT y tipo de documento CDI) fuera del cálculo", True,
        (f"{len(exc)} movimiento(s) excluidos (NIT {', '.join(nits or [])} o tipo de documento {', '.join(TIPOS_DOC_EXCLUIDOS)}) de la base y de las retenciones (ver detalle)."
         if exc else f"Sin movimientos excluidos (NIT: {', '.join(nits or []) or 'ninguno configurado'})."))
    ret = retenciones or []
    add("Retenciones 2365 / 2368 incluidas en el informe", bool(ret),
        f"{len(ret)} cuenta(s) de retención en el período." if ret else
        "No hay movimientos en cuentas 2365 ni 2368 para la empresa y el período.", "warn")
    contab = sum((c.creditos for c in ret if c.tipo == "autorretencion"), Decimal(0))
    if auto is not None and contab:
        dif = contab - auto
        add("Autorretención contabilizada (cuenta 2365 'AUTORRETENCION') vs. calculada", abs(dif) < Decimal("1"),
            f"Contabilizada: {fmt(contab)} · Calculada: {fmt(auto)} · Diferencia: {fmt(dif)}.", "warn")
    return out


# ================================================================= MÓDULO IVA (cuenta 2408 + 135517 + clase 4)
PREFIJO_IVA = "2408"
SIN_TARIFA = "Sin tarifa"

# Clasificación de las cuentas 2408 según el NOMBRE de la cuenta (la devolución tiene prioridad):
#   dev_ventas  → "…devolución…ventas…"   (naturaleza débito: reduce el IVA generado)
#   dev_compras → "…devolución…compras…"  (naturaleza crédito: reduce el IVA descontable)
#   descontable → "…descontable…"         (naturaleza débito)
#   generado    → "…generado…"            (naturaleza crédito)
#   otro / dev_sin_clasificar             (naturaleza crédito; se muestran aparte y se advierte)
NATURALEZA_DEBITO = {"descontable", "dev_ventas"}
TIPOS_IVA = {"generado": "IVA generado", "dev_ventas": "Devolución en ventas", "descontable": "IVA descontable",
             "dev_compras": "Devolución en compras", "dev_sin_clasificar": "Devolución sin clasificar (¿compras o ventas?)",
             "otro": "Otro"}


def es_cuenta_iva(codigo: str) -> bool:
    return str(codigo).strip().startswith(PREFIJO_IVA)


def _sin_acentos(s: str) -> str:
    return "".join(ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn")


def tipo_iva(nombre: str) -> str:
    u = _sin_acentos(nombre).upper()
    if "DEVOL" in u:
        if "COMPRA" in u:
            return "dev_compras"
        if "VENTA" in u:
            return "dev_ventas"
        return "dev_sin_clasificar"
    if "DESCONT" in u:
        return "descontable"
    if "GENERAD" in u:
        return "generado"
    return "otro"


@dataclass
class CuentaIVA:
    codigo: str
    nombre: str
    tipo: str                                   # ver TIPOS_IVA
    creditos: Decimal = Decimal(0)              # movimientos con valor < 0
    debitos: Decimal = Decimal(0)               # movimientos con valor > 0
    base_cred: Decimal = Decimal(0)
    base_deb: Decimal = Decimal(0)
    por_tarifa: dict = field(default_factory=dict)   # tarifa -> [base_cred, créditos, base_deb, débitos]
    movs: list[dict] = field(default_factory=list)

    @property
    def natural(self) -> Decimal:
        """Valor neto en la naturaleza de la cuenta (positivo = comportamiento normal)."""
        return (self.debitos - self.creditos) if self.tipo in NATURALEZA_DEBITO else (self.creditos - self.debitos)

    @property
    def n(self) -> int:
        return len(self.movs)


def build_iva(movs: Iterable[dict]) -> list[CuentaIVA]:
    acc: dict[str, CuentaIVA] = {}
    for m in movs:
        if not es_cuenta_iva(m["codigocuenta"]):
            continue
        c = acc.setdefault(m["codigocuenta"], CuentaIVA(m["codigocuenta"], m["nombrecuenta"], tipo_iva(m["nombrecuenta"])))
        c.movs.append(m)
        v, base = m["valor"], m.get("baseimpuesto", Decimal(0))
        pct = m.get("porcentajeimpuesto") or Decimal(0)
        p = c.por_tarifa.setdefault(pct.normalize() if pct else SIN_TARIFA, [Decimal(0)] * 4)
        if v < 0:
            c.creditos += -v; c.base_cred += base; p[0] += base; p[1] += -v
        elif v > 0:
            c.debitos += v; c.base_deb += base; p[2] += base; p[3] += v
    return sorted(acc.values(), key=lambda c: c.codigo)


def _natural_tarifa(c: CuentaIVA, v: list) -> Decimal:
    return (v[3] - v[1]) if c.tipo in NATURALEZA_DEBITO else (v[1] - v[3])


def totales_iva(cuentas: list[CuentaIVA]) -> dict:
    z = Decimal(0)
    s = lambda tipos: sum((c.natural for c in cuentas if c.tipo in tipos), z)  # noqa: E731
    gen, dv = s({"generado"}), s({"dev_ventas"})
    des, dc = s({"descontable"}), s({"dev_compras"})
    otros = s({"otro", "dev_sin_clasificar"})
    gen_neto, des_neto = gen - dv, des - dc
    return dict(generado=gen, dev_ventas=dv, generado_neto=gen_neto, descontable=des, dev_compras=dc, descontable_neto=des_neto,
                otros=otros, neto_2408=gen_neto - des_neto + otros,
                creditos=sum((c.creditos for c in cuentas), z), debitos=sum((c.debitos for c in cuentas), z),
                base_cred=sum((c.base_cred for c in cuentas), z), base_deb=sum((c.base_deb for c in cuentas), z))


def iva_por_tarifa(cuentas: list[CuentaIVA]) -> list[dict]:
    """Clasificación por tarifa: IVA generado, devoluciones en ventas, descontable, devoluciones en compras, otros y base."""
    agg: dict = {}
    for c in cuentas:
        for t, v in c.por_tarifa.items():
            a = agg.setdefault(t, dict(base_gen=Decimal(0), generado=Decimal(0), dev_ventas=Decimal(0), base_des=Decimal(0),
                                       descontable=Decimal(0), dev_compras=Decimal(0), otros=Decimal(0)))
            n = _natural_tarifa(c, v)
            if c.tipo == "generado":
                a["generado"] += n; a["base_gen"] += v[0]
            elif c.tipo == "dev_ventas":
                a["dev_ventas"] += n
            elif c.tipo == "descontable":
                a["descontable"] += n; a["base_des"] += v[2]
            elif c.tipo == "dev_compras":
                a["dev_compras"] += n
            else:
                a["otros"] += n
    key = lambda t: (t == SIN_TARIFA, t if t != SIN_TARIFA else Decimal(0))  # noqa: E731
    out = []
    for t, a in sorted(agg.items(), key=lambda kv: key(kv[0])):
        out.append(dict(tarifa=t, **a, neto=a["generado"] - a["dev_ventas"] - (a["descontable"] - a["dev_compras"]) + a["otros"]))
    return out


# ---- 135517: IVA retenido (retenciones de IVA que le aplican a la empresa; cuenta de naturaleza débito)
PREFIJO_RETIVA = "135517"


def es_cuenta_reteiva(codigo: str) -> bool:
    return str(codigo).strip().startswith(PREFIJO_RETIVA)


@dataclass
class CuentaRetIVA:
    codigo: str
    nombre: str
    aplicada: Decimal = Decimal(0)              # débitos (valor > 0): retención de IVA que le practicaron
    reversion: Decimal = Decimal(0)             # créditos (valor < 0): reversiones, devoluciones, aplicaciones
    base_apl: Decimal = Decimal(0)
    base_rev: Decimal = Decimal(0)
    por_tarifa: dict = field(default_factory=dict)   # tarifa -> [base_apl, aplicada, base_rev, reversion]
    movs: list[dict] = field(default_factory=list)

    @property
    def neto(self) -> Decimal:
        return self.aplicada - self.reversion

    @property
    def n(self) -> int:
        return len(self.movs)


def build_reteiva(movs: Iterable[dict]) -> list[CuentaRetIVA]:
    acc: dict[str, CuentaRetIVA] = {}
    for m in movs:
        if not es_cuenta_reteiva(m["codigocuenta"]):
            continue
        c = acc.setdefault(m["codigocuenta"], CuentaRetIVA(m["codigocuenta"], m["nombrecuenta"]))
        c.movs.append(m)
        v, base = m["valor"], m.get("baseimpuesto", Decimal(0))
        pct = m.get("porcentajeimpuesto") or Decimal(0)
        p = c.por_tarifa.setdefault(pct.normalize() if pct else SIN_TARIFA, [Decimal(0)] * 4)
        if v > 0:
            c.aplicada += v; c.base_apl += base; p[0] += base; p[1] += v
        elif v < 0:
            c.reversion += -v; c.base_rev += base; p[2] += base; p[3] += -v
    return sorted(acc.values(), key=lambda c: c.codigo)


def totales_reteiva(cuentas: list[CuentaRetIVA]) -> dict:
    z = Decimal(0)
    apl, rev = sum((c.aplicada for c in cuentas), z), sum((c.reversion for c in cuentas), z)
    return dict(aplicada=apl, reversion=rev, neto=apl - rev,
                base_apl=sum((c.base_apl for c in cuentas), z), base_rev=sum((c.base_rev for c in cuentas), z))


def liquidar_iva(iva: list[CuentaIVA], reteiva: list[CuentaRetIVA], invertir: bool = True) -> dict:
    """IVA por pagar = IVA generado − devoluciones en ventas − (IVA descontable − devoluciones en compras) ± otros − IVA retenido.
    Cada línea se presenta con su EFECTO sobre el IVA por pagar. invertir=True: +aumenta / −disminuye;
    invertir=False: signo del auxiliar (pasivo negativo, todo con el signo contrario)."""
    t, r = totales_iva(iva), totales_reteiva(reteiva)
    lin = dict(generado=t["generado"], dev_ventas=-t["dev_ventas"], descontable=-t["descontable"], dev_compras=t["dev_compras"],
               otros=t["otros"])
    neto_2408 = sum(lin.values(), Decimal(0))
    por_pagar = neto_2408 - r["neto"]
    sg = 1 if invertir else -1
    return dict(**{k: v * sg for k, v in lin.items()}, generado_neto=t["generado_neto"] * sg, descontable_neto=-t["descontable_neto"] * sg,
                neto_2408=neto_2408 * sg, retenido=-r["neto"] * sg, por_pagar=por_pagar * sg)


def ejecutar_controles_iva(movs_periodo: list[dict], ingresos: Base, iva: list[CuentaIVA], total_movs: int,
                           claves_unicas: int, excluidos: list[dict] | None = None, nits: list[str] | None = None,
                           reteiva: list[CuentaRetIVA] | None = None) -> list[dict]:
    out = []

    def add(titulo, ok, detalle, nivel="error"):
        out.append(dict(titulo=titulo, ok=ok, detalle=detalle, nivel="ok" if ok else nivel))

    t = totales_iva(iva)
    rt = totales_reteiva(reteiva or [])
    add("Todas las cuentas de ingresos comienzan por 4 y las de IVA por 2408", all(es_cuenta_ingreso(c.codigo) for c in ingresos.cuentas)
        and all(es_cuenta_iva(c.codigo) for c in iva),
        f"{len(ingresos.cuentas)} cuenta(s) de ingresos y {len(iva)} cuenta(s) 2408.")
    add("Total de ingresos = suma neta de la columna `valor` de la clase 4", True,
        f"Suma de `valor` (signo del auxiliar): {fmt(ingresos.base_raw)}. No se usan débito, crédito ni saldos.")
    add("Hay movimientos en cuentas 2408 en el período", bool(iva),
        f"{len(iva)} cuenta(s) 2408." if iva else "No hay movimientos 2408 para la empresa y el período.", "warn")
    sin = [c for c in iva if c.tipo in ("dev_sin_clasificar", "otro")]
    add("Todas las cuentas 2408 clasificadas (generado / descontable / devolución en ventas / devolución en compras)", not sin,
        ("Sin clasificar por el nombre: " + "; ".join(f"{c.codigo} {c.nombre}" for c in sin) +
         ". Se tratan como «otros» (naturaleza crédito).") if sin else
        "La clasificación sale del nombre de la cuenta (DEVOLUCIÓN + COMPRAS/VENTAS, DESCONTABLE, GENERADO).", "warn")
    add("Hay movimientos de la empresa en el período", bool(movs_periodo),
        f"{len(movs_periodo)} movimiento(s) en el rango." if movs_periodo else "No hay movimientos en el rango elegido.")
    add("Hay movimientos de IVA retenido (cuenta 135517) en el período", bool(reteiva),
        f"{len(reteiva or [])} cuenta(s) 135517; retención de IVA neta aplicada: {fmt(rt['neto'])}." if reteiva else
        "No hay movimientos en cuentas 135517 (IVA retenido) para la empresa y el período.", "warn")
    pagar = t["neto_2408"] - rt["neto"]
    add("IVA por pagar (generado − devol. ventas − descontable + devol. compras − retenido)", pagar >= 0,
        f"IVA por pagar: {fmt(pagar)}." if pagar >= 0 else f"Resultado negativo: saldo a favor de {fmt(-pagar)}.", "warn")
    contrarias = [c for c in iva if c.natural < 0]
    add("Cuentas 2408 con saldo contrario a su naturaleza", not contrarias,
        ("; ".join(f"{c.codigo} {c.nombre}: {fmt(c.natural)}" for c in contrarias)) if contrarias else
        "Ninguna cuenta 2408 con saldo contrario a su naturaleza.", "warn")
    ing = ingresos.base
    if ing and iva and ingresos.invertir_signo:
        ef = t["generado_neto"] / ing * 100
        add("IVA generado (neto de devoluciones) frente al total de ingresos", ef <= 19,
            f"IVA generado neto {fmt(t['generado_neto'])} / ingresos {fmt(ing)} = {fmt(ef)}%"
            + ("" if ef <= 19 else " — supera la tarifa general (19%): revise clasificación o exclusiones."), "warn")
    add("Sin movimientos duplicados", total_movs == claves_unicas,
        f"{total_movs} movimiento(s) únicos almacenados (clave: empresa + id del auxiliar).")
    exc = excluidos or []
    add("Movimientos excluidos (NIT y tipo de documento CDI) fuera del cálculo", True,
        (f"{len(exc)} movimiento(s) excluidos (NIT {', '.join(nits or [])} o tipo de documento {', '.join(TIPOS_DOC_EXCLUIDOS)}) (ver detalle)." if exc
         else f"Sin movimientos excluidos (NIT: {', '.join(nits or []) or 'ninguno configurado'})."))
    return out
