"""Exportación a Excel del módulo de impuesto sobre las ventas (IVA)."""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from core import Base, CuentaIVA, CuentaRetIVA, iva_por_tarifa, liquidar_iva, totales_iva, totales_reteiva
from excel import NUM, _f, _sheet

TIPO = {"generado": "IVA generado", "descontable": "IVA descontable", "otro": "Otro"}


def _negrita(ws, desde_fila=2, marcas=("TOTAL",)):
    for row in ws.iter_rows(min_row=desde_fila):
        if str(row[0].value or "").startswith(marcas):
            for c in row:
                c.font = Font(bold=True)


def exportar_iva(empresa: str, desde: str, hasta: str, ingresos: Base, iva: list[CuentaIVA], checks: list[dict],
                 excluidos: list[dict] | None = None, nits: list[str] | None = None,
                 reteiva: list[CuentaRetIVA] | None = None) -> bytes:
    exc = excluidos or []
    rv = reteiva or []
    t = totales_iva(iva)
    rt = totales_reteiva(rv)
    inv = ingresos.invertir_signo
    liq = liquidar_iva(iva, rv, inv)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumen")
    filas = [
        ["IMPUESTO SOBRE LAS VENTAS (IVA) — liquidación preliminar sujeta a revisión"], [],
        ["Empresa", empresa], ["Período", f"{desde} a {hasta}"],
        ["Signo de presentación", "Invertido (ingresos e IVA generado positivos)" if inv else "Tal cual el auxiliar"],
        ["NIT excluidos del cálculo", ", ".join(nits or []) or "Ninguno"], ["Movimientos excluidos por NIT", len(exc)], [],
        ["Total ingresos (cuentas clase 4, neto)", _f(ingresos.base)],
        ["IVA generado (2408, créditos)", _f(liq["generado"])],
        ["(−) IVA descontable / devoluciones / otros débitos (2408)", _f(liq["descontable"])],
        ["IVA neto 2408 (generado − descontable)", _f(liq["neto_2408"])],
        ["(−) IVA retenido (cuenta 135517, neto de reversiones)", _f(liq["retenido"])],
        ["IVA POR PAGAR (generado − descontable − retenido)", _f(liq["por_pagar"])],
        ["IVA generado / total ingresos (%)", float(t["generado"] / ingresos.base * 100) if ingresos.base and inv else None], [],
        ["Esta liquidación no sustituye la validación de las reglas tributarias aplicables a la empresa y al período."],
    ]
    for f in filas:
        ws.append(f)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 58, 46
    for r in range(9, 15):
        ws[f"B{r}"].number_format = NUM
    for r in (9, 12, 14):
        ws[f"A{r}"].font = Font(bold=True)
    ws["B15"].number_format = "0.00"

    ing = [["Código", "Cuenta", "Movimientos", "Suma valor (auxiliar)", "Valor neto de la cuenta"]]
    ing += [[c.codigo, c.nombre, c.n, c.raw, c.valor] for c in ingresos.cuentas]
    ing.append(["TOTAL INGRESOS (clase 4)", "", ingresos.n_movs, ingresos.base_raw, ingresos.base])
    _negrita(_sheet(wb, "Ingresos clase 4", ing, [14, 44, 12, 24, 24], "DE"))

    sg = 1 if inv else -1
    cu = [["Código", "Cuenta", "Clasificación", "Tarifa(s) %", "Base (créditos)", "IVA generado (créditos)",
           "Base (débitos)", "Descontable / devoluciones / débitos (−)", "Neto", "Movimientos"]]
    for c in iva:
        cu.append([c.codigo, c.nombre, TIPO[c.tipo], ", ".join(f"{k:f}" if k != "Sin tarifa" else k for k in c.por_tarifa),
                   c.base_cred, c.generado, c.base_deb, -c.debitos, c.neto, c.n])
    cu.append(["TOTAL IVA 2408", "", "", "", t["base_cred"], t["generado"], t["base_deb"], -t["debitos"], t["generado"] - t["debitos"], ""])
    _negrita(_sheet(wb, "IVA 2408 por cuenta", cu, [14, 44, 18, 14, 22, 22, 22, 30, 20, 12], "EFGHI"))

    tf = [["Tarifa %", "Base (créditos)", "IVA generado", "Base (débitos)", "Descontable / devoluciones / débitos (−)", "Neto"]]
    for r in iva_por_tarifa(iva):
        tf.append([f"{r['tarifa']:f}" if r["tarifa"] != "Sin tarifa" else r["tarifa"], r["base_cred"], r["generado"], r["base_deb"],
                   -r["debitos"], r["neto"]])
    tf.append(["TOTAL", t["base_cred"], t["generado"], t["base_deb"], -t["debitos"], t["generado"] - t["debitos"]])
    _negrita(_sheet(wb, "IVA por tarifa", tf, [14, 22, 22, 22, 36, 20], "BCDEF"))

    ri = [["Código", "Cuenta", "Tarifa(s) %", "Base (aplicadas)", "Retención aplicada (débitos)", "Reversiones (−)", "Neto", "Movimientos"]]
    for c in rv:
        ri.append([c.codigo, c.nombre, ", ".join(f"{k:f}" if k != "Sin tarifa" else k for k in c.por_tarifa), c.base_apl, c.aplicada,
                   -c.reversion, c.neto, c.n])
    ri.append(["TOTAL IVA RETENIDO (135517)", "", "", rt["base_apl"], rt["aplicada"], -rt["reversion"], rt["neto"], ""])
    _negrita(_sheet(wb, "IVA retenido 135517", ri, [14, 44, 14, 22, 26, 20, 20, 12], "DEFG"))

    mi = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación", "Tercero",
           "Centro costo", "Descripción", "Base impuesto", "% impuesto", "valor"]]
    for c in iva:
        mi += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
                m["identificacion"], m["tercero"], m["centrocosto"], m["descripcion"], m.get("baseimpuesto"),
                m.get("porcentajeimpuesto"), m["valor"]] for m in c.movs]
    _sheet(wb, "Mov. IVA 2408", mi, [12, 30, 10, 10, 12, 8, 10, 14, 30, 20, 40, 18, 10, 18], "LN")

    mr = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación", "Tercero",
           "Descripción", "Base impuesto", "% impuesto", "valor"]]
    for c in rv:
        mr += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
                m["identificacion"], m["tercero"], m["descripcion"], m.get("baseimpuesto"), m.get("porcentajeimpuesto"), m["valor"]] for m in c.movs]
    _sheet(wb, "Mov. IVA retenido", mr, [12, 30, 10, 10, 12, 8, 10, 14, 30, 40, 18, 10, 18], "KM")

    mv = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación", "Tercero",
           "Descripción", "Centro costo", "valor"]]
    for c in ingresos.cuentas:
        mv += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
                m["identificacion"], m["tercero"], m["descripcion"], m["centrocosto"], m["valor"]] for m in c.movs]
    _sheet(wb, "Mov. ingresos", mv, [12, 30, 10, 10, 12, 8, 10, 14, 30, 40, 20, 18], "L")

    ex = [["Cuenta", "Nombre cuenta", "Fila archivo", "Id auxiliar", "Fecha", "Tipo doc.", "Nro doc.", "Identificación", "Tercero",
           "Descripción", "valor"]]
    ex += [[m["codigocuenta"], m["nombrecuenta"], m["fila"], m["id"], m["fecha"], m["tipodocumento"], m["nrodocumento"],
            m["identificacion"], m["tercero"], m["descripcion"], m["valor"]] for m in exc]
    _sheet(wb, "Excluidos por NIT", ex, [12, 40, 10, 10, 12, 8, 10, 14, 24, 60, 18], "K")

    vl = [["Control", "Resultado", "Detalle"]] + [[c["titulo"], "OK" if c["ok"] else ("Revisar" if c["nivel"] == "warn" else "Falla"), c["detalle"]] for c in checks]
    _sheet(wb, "Validaciones", vl, [56, 10, 100])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
